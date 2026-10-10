from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.export.base import BaseExporter


class ExcelExporter(BaseExporter):
    extension = "xlsx"

    def export(self, data: list[dict], out_dir: Path) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Audit"

        ws.append(["Nhóm", "Số file", "Danh sách file", "Báo cáo"])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for row in data:
            ws.append([
                row.get("chunk"),
                len(row.get("files", [])),
                "\n".join(row.get("files", [])),
                row.get("report", ""),
            ])

        for col, width in zip("ABCD", (8, 10, 40, 100)):
            ws.column_dimensions[col].width = width
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        path = self._make_path(out_dir)
        wb.save(path)
        return path
