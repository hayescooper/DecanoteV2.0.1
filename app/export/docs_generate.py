from pathlib import Path
import re

from docx import Document
from docx.shared import Inches

from app.export.base import BaseExporter


class DocsExporter(BaseExporter):
    extension = "docx"

    def export(self, data: list[dict], out_dir: Path) -> Path:
        doc = Document()
        doc.add_heading("Báo cáo audit - Decanote", level=0)

        for row in data:
            files = row.get("files", [])
            doc.add_heading(f"Nhóm {row.get('chunk')} ({len(files)} file)", level=1)
            
            # Create a structured table for files
            if files:
                table = doc.add_table(rows=1, cols=2)
                table.style = 'Table Grid'
                
                # Header row
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = "STT"
                hdr_cells[1].text = "Đường dẫn File"
                
                # Populate rows
                for idx, name in enumerate(files, start=1):
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(idx)
                    row_cells[1].text = name
                
                # Add spacing after the table
                doc.add_paragraph()

            # Parse report text and convert markdown elements cleanly
            report = row.get("report", "")
            if report:
                for line in report.splitlines():
                    stripped = line.strip()
                    if not stripped:
                        continue
                    
                    if stripped.startswith("### "):
                        doc.add_heading(stripped[4:].strip(), level=3)
                    elif stripped.startswith("## "):
                        doc.add_heading(stripped[3:].strip(), level=2)
                    elif stripped.startswith("# "):
                        doc.add_heading(stripped[2:].strip(), level=1)
                    elif stripped.startswith(("* ", "- ")):
                        # Convert markdown bullet points into native Word bullets
                        bullet_text = re.sub(r'^\s*[\*\-]\s+', '', stripped)
                        self._add_markdown_paragraph(doc, bullet_text, style='List Bullet')
                    elif stripped.startswith("`") and stripped.endswith("`"):
                        # Format standalone code lines as indented code blocks
                        code_text = stripped[1:-1]
                        p = doc.add_paragraph(style='Normal')
                        p.paragraph_format.left_indent = Inches(0.25)
                        run = p.add_run(code_text)
                        run.font.name = 'Consolas'
                    else:
                        self._add_markdown_paragraph(doc, stripped)

        path = self._make_path(out_dir)
        doc.save(path)
        return path

    def _add_markdown_paragraph(self, doc, text, style=None):
        """Helper to parse inline bold (**text**) and code (`text`) inside paragraphs."""
        p = doc.add_paragraph(style=style)
        tokens = re.split(r'(\*\*.*?\*\*|`.*?\`)', text)
        for token in tokens:
            if not token:
                continue
            if token.startswith('**') and token.endswith('**'):
                run = p.add_run(token[2:-2])
                run.bold = True
            elif token.startswith('`') and token.endswith('`'):
                run = p.add_run(token[1:-1])
                run.font.name = 'Consolas'
            else:
                p.add_run(token)
        return p